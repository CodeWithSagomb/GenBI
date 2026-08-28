"""Génère RuwaGenBI_presentation.pptx à partir du contenu du diaporama HTML.

Usage : python3 build_pptx.py
Ne dépend d'aucune donnée dynamique du projet, contenu recopié à la main
depuis docs/presentation/index.html pour rester synchronisé.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK = RGBColor(0x16, 0x32, 0x3E)
PAPER = RGBColor(0xEE, 0xF0, 0xE9)
ACCENT = RGBColor(0xB9, 0x6E, 0x23)
ACCENT_SOFT = RGBColor(0xE4, 0xB2, 0x7C)
TEAL = RGBColor(0x2C, 0x6D, 0x62)
TEXT_DARK = RGBColor(0x1C, 0x25, 0x21)
TEXT_SOFT = RGBColor(0x52, 0x58, 0x4E)
PAPER_TEXT = RGBColor(0xE8, 0xE6, 0xDC)

TITLE_FONT = "Georgia"
BODY_FONT = "Avenir Next"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(dark=False):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK if dark else PAPER
    return slide


def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def set_run(run, text, size, color, bold=False, italic=False, font=BODY_FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def eyebrow(slide, text, dark=False, color=None):
    tf = add_textbox(slide, 0.9, 0.55, 8, 0.5)
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, text.upper(), 13, color or (ACCENT_SOFT if dark else ACCENT), bold=True)
    r.font.name = BODY_FONT
    return tf


def title(slide, text, dark=False, top=1.05, size=34):
    tf = add_textbox(slide, 0.9, top, 11.4, 1.6)
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, text, size, PAPER_TEXT if dark else INK, bold=True, font=TITLE_FONT)
    return tf


def bullets(slide, items, top=2.7, left=0.9, width=11.4, size=18, dark=False, dash_color=None):
    tf = add_textbox(slide, left, top, width, 4.2)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r = p.add_run()
        set_run(r, "•  " + item, size, PAPER_TEXT if dark else TEXT_DARK)
        r.font.color.rgb = (dash_color or ACCENT) if False else r.font.color.rgb
    return tf


def note(slide, text, top, dark=False):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(top), Inches(11.4), Inches(1.1))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x24, 0x41, 0x3A) if dark else RGBColor(0xE4, 0xE7, 0xDD)
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, text, 15, PAPER_TEXT if dark else TEXT_DARK, italic=True, font=TITLE_FONT)
    return box


def set_notes(slide, text):
    """Notes de présentateur : visibles en mode Présentateur, jamais projetées."""
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.text = text


def counter(slide, n, total=18, dark=False):
    tf = add_textbox(slide, 12.35, 7.05, 0.9, 0.35)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    set_run(r, f"{n:02d} / {total}", 10, TEXT_SOFT if not dark else RGBColor(0x9B, 0xA3, 0x9A))


# 1. TITLE ------------------------------------------------------------
s = add_slide(dark=True)
tf = add_textbox(s, 0.9, 2.3, 10, 0.5)
r = tf.paragraphs[0].add_run()
set_run(r, "PRÉSENTATION TECHNIQUE · END-TO-END", 13, ACCENT_SOFT, bold=True)
tf = add_textbox(s, 0.85, 2.75, 11, 1.6)
r = tf.paragraphs[0].add_run()
set_run(r, "RuwaGenBI", 60, PAPER_TEXT, bold=True, font=TITLE_FONT)
tf = add_textbox(s, 0.9, 4.15, 9, 1.2)
p = tf.paragraphs[0]
r = p.add_run()
set_run(r, "BI générative pour pharmacies à Dakar. Du langage naturel au SQL fiable, entièrement en local.", 20, PAPER_TEXT)
tf = add_textbox(s, 0.9, 6.3, 11, 0.5)
r = tf.paragraphs[0].add_run()
set_run(r, "FastAPI · React · Ollama · Postgres · dbt        Zéro fuite de données", 12, RGBColor(0xB8, 0xB6, 0xAC))
set_notes(s, (
    "0:30\n"
    "Accroche : un pharmacien tape une question en français, neuf étapes plus tard il a un "
    "graphique et une phrase, sans qu'aucune donnée n'ait quitté la machine, et sans qu'un "
    "développeur ait écrit la requête à l'avance.\n"
    "Se présenter brièvement, planter le décor (pharmacies de Dakar)."
))
counter(s, 1, dark=True)

# 2. LE PROBLEME --------------------------------------------------------
s = add_slide()
eyebrow(s, "Le contexte")
title(s, "La BI classique fige la question avant de connaître le besoin")
bullets(s, [
    "Besoin métier, puis analyste qui écrit le SQL, puis dashboard publié, puis l'utilisateur consulte.",
    "Toute question non anticipée déclenche un nouveau cycle : ticket, développeur, déploiement.",
    "L'utilisateur reste prisonnier de ce qui a été construit à l'avance.",
], top=2.6)
note(s, "Dans ce projet même, le Dashboard React interroge 6 requêtes SQL pré-écrites. C'est de la BI classique, avec une interface moderne.", top=5.6)
set_notes(s, (
    "1:00\n"
    "Dérouler le cycle dessiné à l'écran : besoin, analyste, dashboard, consultation.\n"
    "Insister : toute question non anticipée = un ticket, un cycle de développement complet.\n"
    "Exemple concret dans ce projet même : le Dashboard React fait exactement ça, SQL pré-écrit."
))
counter(s, 2)

# 3. LA PROPOSITION ------------------------------------------------------
s = add_slide()
eyebrow(s, "La proposition")
title(s, "Générer le SQL à la volée, pas l'écrire à l'avance")
bullets(s, [
    "Un LLM traduit n'importe quelle question en langage naturel vers du SQL, au moment où elle est posée.",
    "La couverture n'est plus ce qu'on a eu le temps de construire, mais tout ce que les données permettent de répondre.",
], top=2.7)
note(s, "Le vrai piège : un LLM peut halluciner du SQL faux. Le reste de la présentation, c'est comment on construit la confiance autour de cette imprévisibilité.", top=5.2)
set_notes(s, (
    "1:00\n"
    "Poser la proposition centrale : générer plutôt qu'écrire à l'avance.\n"
    "Nommer tout de suite le risque, l'hallucination, pour désamorcer la question qui vient.\n"
    "Annoncer explicitement : le reste de la présentation répond à comment on construit la confiance."
))
counter(s, 3)

# 4. ARCHITECTURE ---------------------------------------------------------
s = add_slide()
eyebrow(s, "Vue d'ensemble")
title(s, "Trois couches, un LLM local, une base gardée par RLS")
bullets(s, [
    "Frontend : React 18 et Vite. Dashboard (SQL fixe), Chat (LLM), Profil. Français et anglais.",
    "Backend : FastAPI et Python 3.11. Orchestration du pipeline en 9 étapes. Aucune logique métier codée en dur.",
    "LLM local : Ollama tourne nativement sur macOS avec qwen2.5-coder:7b, jamais dans un container Docker.",
], top=2.6)
note(s, "Concrètement, aucune question posée ni aucune donnée ne sort jamais de la machine.", top=5.7)
set_notes(s, (
    "1:00\n"
    "Décrire les 3 couches rapidement, une phrase chacune.\n"
    "Insister sur Ollama natif macOS, jamais dans Docker, zéro appel réseau externe.\n"
    "C'est la base de l'argument confidentialité qui revient plus tard (slide sécurité)."
))
counter(s, 4)

# 5. DONNEES ---------------------------------------------------------------
s = add_slide()
eyebrow(s, "La couche de données", color=TEAL)
title(s, "De raw à staging à marts, une isolation qui ne dépend pas du code")
bullets(s, [
    "19 modèles dbt au total : 10 en staging, 9 en marts.",
    "Chaque table de faits reçoit une règle de sécurité au niveau de la ligne : une vente n'est visible que si son identifiant de pharmacie correspond à la session en cours.",
    "Cette règle est appliquée automatiquement à chaque reconstruction des tables, sans intervention manuelle.",
], top=2.6)
note(s, "Même si le SQL généré oublie un filtre, la fuite entre pharmacies reste structurellement impossible.", top=5.7)
set_notes(s, (
    "1:00\n"
    "raw, staging, marts : une phrase chacune, ne pas s'attarder.\n"
    "Point clé : l'isolation est une policy Postgres (RLS), pas un WHERE écrit en Python.\n"
    "Répéter la phrase de garantie : même si le LLM oublie un filtre, impossible de voir une autre pharmacie."
))
counter(s, 5)

# 6. PIPELINE OVERVIEW -------------------------------------------------------
s = add_slide()
eyebrow(s, "Le pipeline")
title(s, "9 étapes, environ 800 ms, zéro donnée qui quitte la machine")
steps = [
    "Recherche d'exemples similaires (ChromaDB)",
    "Filtrage du schéma pertinent (Python pur)",
    "Génération du SQL par le modèle (Ollama, température 0.0)",
    "Validation du SQL, lecture seule (sqlglot)",
    "Exécution, réparation automatique si besoin (Postgres, RLS)",
    "Vérification de la cohérence (Python pur)",
    "Mise en forme lisible (Python pur)",
    "Rédaction de l'insight (Ollama, température 0.1)",
    "Choix automatique du graphique (Python pur)",
]
bullets(s, [f"{i+1}. {t}" for i, t in enumerate(steps)], top=2.5, size=16)
set_notes(s, (
    "1:00\n"
    "Vue d'ensemble seulement, ne pas détailler chaque étape ici.\n"
    "Annoncer qu'on va zoomer sur les points les plus intéressants, pas les 9 en détail.\n"
    "Donner le chiffre clé : environ 800 ms bout en bout sur une question de référence."
))
counter(s, 6)

# 7. RAG + SCHEMA FILTER -------------------------------------------------
s = add_slide()
eyebrow(s, "Étapes 1 et 2")
title(s, "Réduire intelligemment le contexte, sans LLM")
bullets(s, [
    "Recherche d'exemples : une collection par pharmacie dans ChromaDB, les 3 exemples question et SQL les plus proches en sens.",
    "Filtrage du schéma : un score qui combine un peu de lexical et surtout de la similarité de sens, calculé par table.",
    "Garde les 15 tables les plus pertinentes sur 19, dont 3 toujours incluses par sécurité.",
], top=2.6)
note(s, "Aucun appel au LLM dans ce filtrage : un calcul simple, tenu en mémoire, recalculé une seule fois au démarrage.", top=5.7)
set_notes(s, (
    "1:30\n"
    "Deux mécanismes bien distincts, insister sur la différence : exemples passés vs tables pertinentes.\n"
    "Répéter : zéro appel LLM dans le filtrage de schéma, rapide et déterministe.\n"
    "Chiffre à donner : schéma réduit de 13 200 à 2 600 caractères, gain mesuré sur le taux de succès."
))
counter(s, 7)

# 8. GEN SQL + VALIDATE ---------------------------------------------------
s = add_slide()
eyebrow(s, "Étapes 3 et 4")
title(s, "Génération déterministe, validation structurelle")
bullets(s, [
    "Le SQL est généré à température zéro, pour rester le plus déterministe possible.",
    "Les rappels les plus importants sont ajoutés juste avant la question elle-même, le modèle les lit en dernier et les suit mieux.",
    "La validation est un vrai parseur : une seule instruction, et bien une lecture, jamais une écriture.",
], top=2.6)
note(s, "Ce n'est pas là que se joue la vraie sécurité. Elle vient des droits en lecture seule sur la base et de l'isolation au niveau des lignes.", top=5.7)
set_notes(s, (
    "1:30\n"
    "Génération à température 0, expliquer pourquoi : déterminisme, la même question doit produire le même SQL.\n"
    "Détail à mentionner si public curieux : les rappels critiques sont injectés juste avant la question, "
    "position de haute récence, le modèle les lit en dernier et les suit mieux que les règles en tête de prompt.\n"
    "Validation = un vrai parseur (sqlglot), pas une liste de mots interdits.\n"
    "Dire explicitement : ce n'est pas la vraie protection de sécurité, ça vient juste après."
))
counter(s, 8)

# 9. MARS-SQL --------------------------------------------------------------
s = add_slide()
eyebrow(s, "Étape 5, moment fort")
title(s, "MARS-SQL, la réparation automatique par feedback d'exécution")
bullets(s, [
    "Le SQL échoue.",
    "L'erreur exacte remonte au LLM.",
    "Le SQL est corrigé.",
    "Nouvel essai, jusqu'à 3 au total.",
], top=2.6)
note(s, "La différence avec un simple nouvel essai : le modèle voit pourquoi ça a échoué, pas seulement qu'il faut réessayer à l'aveugle.", top=5.5)
set_notes(s, (
    "1:30, MOMENT FORT, prendre le temps ici\n"
    "Dérouler le cycle lentement : le SQL échoue, l'erreur exacte remonte au LLM, il corrige, nouvel essai.\n"
    "Insister sur la différence avec un retry aveugle : le modèle voit pourquoi, pas juste qu'il faut réessayer.\n"
    "Chiffre : jusqu'à 3 tentatives au total (1 initiale plus 2 réparations).\n"
    "Question probable à anticiper : et si ça échoue les 3 fois ? Réponse : une erreur claire est renvoyée à "
    "l'utilisateur, jamais un silence ou un résultat inventé."
))
counter(s, 9)

# 10. SEMANTIC VALIDATION ---------------------------------------------------
s = add_slide()
eyebrow(s, "Étape 6", color=TEAL)
title(s, "La cohérence du résultat, pas seulement sa syntaxe")
bullets(s, [
    "Une réponse vide est suspecte, sauf pour une question d'existence.",
    "Une question qui attend une seule valeur mais reçoit plus de 3 lignes sans regroupement, c'est suspect.",
    "Un top N qui renvoie beaucoup plus que N lignes, c'est suspect aussi.",
], top=2.6)
note(s, "Volontairement conservateur : préfère laisser passer un cas limite plutôt que re-générer inutilement.", top=5.7)
set_notes(s, (
    "1:00\n"
    "Bien distinguer de l'étape précédente : ici le SQL s'exécute sans erreur, la question est différente.\n"
    "Donner les 3 règles rapidement, sans s'attarder sur chacune.\n"
    "Insister : volontairement conservateur, préfère laisser passer un cas limite plutôt que re-générer à tort.\n"
    "C'est la couche que la plupart des systèmes text-to-SQL n'ont pas, à souligner."
))
counter(s, 10)

# 11. HUMANIZE + INSIGHT ----------------------------------------------------
s = add_slide()
eyebrow(s, "Étapes 7 et 8")
title(s, "Du solide construit autour de l'incertain")
bullets(s, [
    "Mise en forme lisible : mois, jours, valeurs vrai ou faux. De simples conversions, sans dépendance au LLM.",
    "L'insight : une phrase générée à faible température, puis corrigée automatiquement pour que le format des montants soit toujours correct.",
], top=2.6)
note(s, "Chaque étape ajoute une sécurité fiable autour d'un modèle qui ne l'est pas toujours. L'architecture part du principe que le LLM va se tromper.", top=5.2)
set_notes(s, (
    "1:00\n"
    "Humanisation : conversions Python pures, zéro dépendance au LLM, aller vite ici.\n"
    "Insight : température faible, puis correction déterministe après coup sur le format des montants.\n"
    "Message clé à répéter mot pour mot si possible : l'architecture part du principe que le LLM va se tromper, "
    "et s'organise autour de ça."
))
counter(s, 11)

# 12. VIZ HINT ---------------------------------------------------------------
s = add_slide()
eyebrow(s, "Étape 9")
title(s, "Le graphique se choisit tout seul")
bullets(s, [
    "Tendance dans le temps → courbe",
    "Classement limité → barres",
    "Comparaison → barres",
    "Répartition → camembert",
    "Sinon → barres, par défaut",
], top=2.6)
note(s, "Une cascade de règles classées par priorité, entièrement calculée à l'avance, sans passer par le LLM.", top=5.9)
set_notes(s, (
    "0:45\n"
    "Dérouler la cascade rapidement, une seule phrase par règle.\n"
    "Répéter : toute cette logique est en Python et en regex, encore zéro LLM ici.\n"
    "Mentionner que le frontend affine encore le choix ensuite (courbe simple ou combinée)."
))
counter(s, 12)

# 13. CONFIG DRIVEN -----------------------------------------------------------
s = add_slide()
eyebrow(s, "Architecture", color=TEAL)
title(s, "Zéro redéploiement pour ajuster le comportement du LLM")
bullets(s, [
    "Toutes les règles métier vivent dans quatre fichiers de configuration lisibles.",
    "Un simple appel à un point d'administration les recharge à chaud, sans redémarrer quoi que ce soit.",
    "Ajuster une règle revient à éditer un fichier texte, jamais à toucher au code Python.",
], top=2.6)
set_notes(s, (
    "1:00\n"
    "Les règles métier vivent en configuration, pas en code, insister sur ce choix.\n"
    "Rechargement à chaud, sans redéploiement : bon exemple pour montrer l'agilité du projet.\n"
    "Si public technique curieux : chaque règle est un fichier lisible, pas une ligne de code cachée."
))
counter(s, 13)

# 14. QUALITY -----------------------------------------------------------------
s = add_slide()
eyebrow(s, "Discipline qualité")
title(s, "Benchmarker au lieu de supposer")
bullets(s, [
    "50 questions métier réparties en 7 catégories, un test rejoué à chaque changement.",
    "Face à un modèle deux fois plus gros : même score, mais deux fois plus lent. Retour en arrière assumé, chiffres à l'appui.",
    "Le modèle retenu (7 milliards de paramètres) offre le meilleur équilibre entre qualité et rapidité, mesuré et non supposé.",
], top=2.6)
set_notes(s, (
    "1:00\n"
    "Golden set de 50 questions, reproductible à chaque changement, à mentionner brièvement.\n"
    "Point fort à vraiment appuyer ici : rejet argumenté d'un modèle plus gros, avec des chiffres précis.\n"
    "Ça montre une discipline d'ingénieur, pas juste on a choisi le plus gros modèle disponible."
))
counter(s, 14)

# 15. SECURITY -----------------------------------------------------------------
s = add_slide()
eyebrow(s, "Sécurité", color=TEAL)
title(s, "Trois couches indépendantes, aucune ne dépend du LLM")
bullets(s, [
    "Un rôle en lecture seule : le compte utilisé par le pipeline n'a que des droits de lecture, et ne peut pas contourner l'isolation.",
    "L'isolation par la base : chaque pharmacie ne voit que ses propres lignes.",
    "Un vrai parseur : une seule instruction, jamais autre chose qu'une lecture.",
], top=2.6)
set_notes(s, (
    "1:00\n"
    "Récapituler les 3 couches, une phrase chacune, ne pas re-détailler ce qui a déjà été dit.\n"
    "Répéter le point le plus important : aucune de ces 3 couches ne dépend de la fiabilité du LLM.\n"
    "C'est la meilleure réponse si quelqu'un demande et si le modèle se trompe complètement."
))
counter(s, 15)

# 16. DEMO -----------------------------------------------------------------
s = add_slide(dark=True)
tf = add_textbox(s, 0.9, 2.3, 10, 0.5)
r = tf.paragraphs[0].add_run()
set_run(r, "MOMENT DÉMO", 13, ACCENT_SOFT, bold=True)
title(s, "Une vraie question, en direct", dark=True, top=2.75, size=40)
tf = add_textbox(s, 0.9, 4.2, 10.5, 1.5)
r = tf.paragraphs[0].add_run()
set_run(r, "On pose une question dans le Chat, et on regarde ensemble le SQL généré, le graphique, et la phrase qui en sort. Si possible, une question un peu ambiguë pour voir la réparation automatique se déclencher en direct.", 18, PAPER_TEXT)
set_notes(s, (
    "2:00 à 3:00, BASCULER SUR L'APPLICATION RÉELLE\n"
    "Poser d'abord une question simple pour montrer le chemin normal (SQL, graphique, insight).\n"
    "Si le temps permet : poser une question un peu ambiguë pour déclencher MARS-SQL en direct.\n"
    "Filet de sécurité prévu : capture d'écran ou vidéo courte de secours si Ollama est lent ce jour-là."
))
counter(s, 16, dark=True)

# 17. LIMITS -----------------------------------------------------------------
s = add_slide()
eyebrow(s, "Honnêteté technique")
title(s, "Ce qui ne scale pas encore")
bullets(s, [
    "Ollama ne traite qu'une génération à la fois. Pas encore de vrai parallélisme entre plusieurs utilisateurs.",
    "Assumé : ce système reste un MVP de labo, pas une mise en production.",
    "Pistes : activer le parallélisme d'Ollama, mettre en cache les questions fréquentes, passer à un serveur d'inférence avec traitement par lots.",
], top=2.6)
set_notes(s, (
    "1:00, DEUXIÈME MOMENT OÙ LE PUBLIC TECHNIQUE VA CREUSER, prévoir de la marge\n"
    "Assumer complètement : Ollama mono-thread, pas de vrai parallélisme aujourd'hui.\n"
    "Contexte à rappeler : MVP de labo, pas une mise en production, c'est un choix assumé.\n"
    "Si la question du scaling revient : donner les 3 pistes affichées, sans s'engager sur un calendrier."
))
counter(s, 17)

# 18. CLOSING -----------------------------------------------------------------
s = add_slide(dark=True)
tf = add_textbox(s, 0.9, 2.0, 10, 0.5)
r = tf.paragraphs[0].add_run()
set_run(r, "CONCLUSION", 13, ACCENT_SOFT, bold=True)
title(s, "Ne pas choisir entre fiabilité et flexibilité", dark=True, top=2.45, size=34)
tf = add_textbox(s, 0.9, 3.9, 10.8, 2.5)
p = tf.paragraphs[0]
r = p.add_run()
set_run(r, "La BI classique gagne en fiabilité ce qu'elle perd en flexibilité. Un LLM livré tel quel fait l'inverse. Ce système garde la liberté du langage naturel, avec des garde-fous solides à chaque étage, pour retrouver la confiance qu'on aurait dans du SQL écrit à la main.", 19, PAPER_TEXT, italic=True, font=TITLE_FONT)
set_notes(s, (
    "0:45\n"
    "Lire la phrase affichée à l'écran, presque telle quelle, elle porte la conclusion.\n"
    "Remercier, ouvrir sur les questions.\n"
    "Garder en tête les 2 sujets probables : MARS-SQL (slide 9) et la scalabilité (slide 17)."
))
counter(s, 18, dark=True)

prs.save("RuwaGenBI_presentation.pptx")
print("OK: RuwaGenBI_presentation.pptx généré,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
